#!/usr/bin/env python3

from pathlib import Path
from collections import Counter, defaultdict
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
import sys


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

REPORT_PERIOD = "2026-Q4"

DATA_DIR = (
    PROJECT_ROOT
    / "reports"
    / REPORT_PERIOD
    / "data"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "reports"
    / REPORT_PERIOD
    / "STC_OS_Health_Summary.docx"
)


CATEGORY_ORDER = [
    "cpu",
    "memory",
    "filesystem",
    "services",
    "kernel",
    "security",
    "logs",
    "Network",
    "NTP",
    "OS Version",
    "Uptime",
    "SAR",
]

CATEGORY_NAMES = {
    "cpu": "CPU",
    "memory": "Memory",
    "filesystem": "Filesystem",
    "services": "Services",
    "kernel": "Kernel",
    "security": "Security",
    "logs": "Logs",
    "Network": "Network",
    "NTP": "NTP / Time Synchronization",
    "OS Version": "OS Version",
    "Uptime": "Uptime",
    "SAR": "SAR / Historical Performance",
}


# ============================================================
# HELPERS
# ============================================================

def load_json_files():
    """Load all server JSON files."""

    if not DATA_DIR.exists():
        print(f"ERROR: Data directory does not exist:")
        print(DATA_DIR)
        sys.exit(1)

    json_files = sorted(DATA_DIR.glob("*.json"))

    if not json_files:
        print("ERROR: No JSON report files found.")
        print(DATA_DIR)
        sys.exit(1)

    servers = []

    for json_file in json_files:
        try:
            with open(json_file, "r", encoding="utf-8") as f:
                data = json.load(f)

            servers.append(data)

        except Exception as exc:
            print(f"WARNING: Could not read {json_file.name}: {exc}")

    if not servers:
        print("ERROR: No valid JSON reports could be loaded.")
        sys.exit(1)

    return servers


def set_cell_text(cell, text, bold=False):
    """Set table cell text."""

    cell.text = ""

    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))

    run.bold = bold
    run.font.size = Pt(9)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    """Apply background shading to a table cell."""

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def add_table_borders(table):
    """Add simple borders to a table."""

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    borders = tbl_pr.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):

        tag = "w:" + edge

        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B7B7B7")


def add_title(document):

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "STC RHEL OS HEALTH CHECK\n"
        "QUARTERLY SUMMARY REPORT"
    )

    run.bold = True
    run.font.size = Pt(20)

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        f"Reporting Period: {REPORT_PERIOD}"
    )

    run.font.size = Pt(12)


def add_section_heading(document, title, level=1):

    heading = document.add_heading(title, level=level)

    return heading


def add_bullet(document, text):

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.add_run(str(text))


def get_status_counts(servers):

    counts = Counter()

    for server in servers:

        health = server.get("health", {})

        status = str(
            health.get("overall_status", "UNKNOWN")
        ).upper()

        counts[status] += 1

    return counts


def get_total_finding_counts(servers):

    counts = Counter()

    for server in servers:

        health = server.get("health", {})

        counts["CRITICAL"] += int(
            health.get("critical", 0)
        )

        counts["WARNING"] += int(
            health.get("warning", 0)
        )

        counts["INFO"] += int(
            health.get("info", 0)
        )

        counts["OK"] += int(
            health.get("ok", 0)
        )

    return counts


def get_all_findings(servers):

    findings = []

    for server in servers:

        hostname = server.get(
            "server", {}
        ).get(
            "hostname",
            "UNKNOWN"
        )

        for finding in server.get("findings", []):

            finding_copy = dict(finding)

            finding_copy["_hostname"] = hostname

            findings.append(finding_copy)

    return findings


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

def add_executive_summary(document, servers):

    add_section_heading(
        document,
        "1. Executive Summary",
        level=1
    )

    status_counts = get_status_counts(servers)

    total_servers = len(servers)

    healthy = status_counts.get("OK", 0)
    warning = status_counts.get("WARNING", 0)
    critical = status_counts.get("CRITICAL", 0)

    paragraph = document.add_paragraph()

    paragraph.add_run(
        f"The quarterly RHEL OS health assessment covered "
        f"{total_servers} server(s) during {REPORT_PERIOD}. "
    )

    paragraph.add_run(
        f"{healthy} server(s) were assessed as healthy, "
        f"{warning} server(s) had warning-level findings, "
        f"and {critical} server(s) had critical findings."
    )

    document.add_paragraph()

    table = document.add_table(
        rows=1,
        cols=4
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    headers = [
        "Status",
        "Servers",
        "Percentage",
        "Meaning",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    for status, meaning in [
        ("OK", "No health concern identified"),
        ("WARNING", "Attention or follow-up recommended"),
        ("CRITICAL", "Immediate investigation required"),
    ]:

        count = status_counts.get(status, 0)

        percentage = (
            (count / total_servers) * 100
            if total_servers
            else 0
        )

        row = table.add_row().cells

        set_cell_text(row[0], status)
        set_cell_text(row[1], count)
        set_cell_text(row[2], f"{percentage:.1f}%")
        set_cell_text(row[3], meaning)

    add_table_borders(table)


# ============================================================
# OVERALL HEALTH SUMMARY
# ============================================================

def add_server_summary(document, servers):

    add_section_heading(
        document,
        "2. Server Health Summary",
        level=1
    )

    table = document.add_table(
        rows=1,
        cols=6
    )

    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
        "Server",
        "OS",
        "Overall",
        "Critical",
        "Warning",
        "Info",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    for server in servers:

        server_info = server.get("server", {})
        health = server.get("health", {})

        row = table.add_row().cells

        hostname = server_info.get(
            "hostname",
            "UNKNOWN"
        )

        os_name = (
            f"{server_info.get('os', 'N/A')} "
            f"{server_info.get('os_version', '')}"
        ).strip()

        status = str(
            health.get(
                "overall_status",
                "UNKNOWN"
            )
        ).upper()

        values = [
            hostname,
            os_name,
            status,
            health.get("critical", 0),
            health.get("warning", 0),
            health.get("info", 0),
        ]

        for i, value in enumerate(values):

            set_cell_text(
                row[i],
                value
            )

    add_table_borders(table)


# ============================================================
# FINDING SUMMARY
# ============================================================

def add_finding_summary(document, servers):

    add_section_heading(
        document,
        "3. Finding Summary",
        level=1
    )

    counts = get_total_finding_counts(servers)

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.style = "Table Grid"

    headers = [
        "Finding Severity",
        "Total Findings",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    for severity in [
        "CRITICAL",
        "WARNING",
        "INFO",
        "OK",
    ]:

        row = table.add_row().cells

        set_cell_text(
            row[0],
            severity
        )

        set_cell_text(
            row[1],
            counts.get(severity, 0)
        )

    add_table_borders(table)

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "Note: Finding counts represent the individual health "
        "checks recorded across all assessed servers. "
        "A single server may contain multiple findings."
    )


# ============================================================
# DOMAIN SUMMARY
# ============================================================

def add_domain_summary(document, servers):

    add_section_heading(
        document,
        "4. Health Domain Summary",
        level=1
    )

    domain_data = defaultdict(
        lambda: Counter()
    )

    for server in servers:

        for finding in server.get("findings", []):

            category = finding.get(
                "category",
                "Unknown"
            )

            severity = str(
                finding.get(
                    "severity",
                    "UNKNOWN"
                )
            ).upper()

            domain_data[category][severity] += 1

    table = document.add_table(
        rows=1,
        cols=5
    )

    table.style = "Table Grid"

    headers = [
        "Health Domain",
        "Critical",
        "Warning",
        "Info",
        "OK",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    categories = list(domain_data.keys())

    ordered_categories = []

    for category in CATEGORY_ORDER:

        if category in categories:

            ordered_categories.append(category)

    for category in categories:

        if category not in ordered_categories:

            ordered_categories.append(category)

    for category in ordered_categories:

        row = table.add_row().cells

        display_name = CATEGORY_NAMES.get(
            category,
            category
        )

        set_cell_text(
            row[0],
            display_name
        )

        for index, severity in enumerate(
            [
                "CRITICAL",
                "WARNING",
                "INFO",
                "OK",
            ],
            start=1
        ):

            set_cell_text(
                row[index],
                domain_data[category].get(
                    severity,
                    0
                )
            )

    add_table_borders(table)


# ============================================================
# CRITICAL FINDINGS
# ============================================================

def add_critical_findings(document, servers):

    add_section_heading(
        document,
        "5. Critical Findings",
        level=1
    )

    findings = [
        finding
        for finding in get_all_findings(servers)
        if str(
            finding.get(
                "severity",
                ""
            )
        ).upper() == "CRITICAL"
    ]

    if not findings:

        paragraph = document.add_paragraph()

        paragraph.add_run(
            "No CRITICAL findings were identified "
            "during the assessment."
        )

        return

    table = document.add_table(
        rows=1,
        cols=4
    )

    table.style = "Table Grid"

    headers = [
        "Server",
        "Domain",
        "Finding",
        "Recommendation",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "F4CCCC"
        )

    for finding in findings:

        row = table.add_row().cells

        set_cell_text(
            row[0],
            finding.get(
                "_hostname",
                "UNKNOWN"
            )
        )

        set_cell_text(
            row[1],
            CATEGORY_NAMES.get(
                finding.get(
                    "category",
                    "Unknown"
                ),
                finding.get(
                    "category",
                    "Unknown"
                )
            )
        )

        set_cell_text(
            row[2],
            finding.get(
                "message",
                "N/A"
            )
        )

        set_cell_text(
            row[3],
            finding.get(
                "recommendation",
                "Review the finding and investigate."
            )
        )

    add_table_borders(table)


# ============================================================
# WARNING FINDINGS
# ============================================================

def add_warning_findings(document, servers):

    add_section_heading(
        document,
        "6. Warning Findings",
        level=1
    )

    findings = [
        finding
        for finding in get_all_findings(servers)
        if str(
            finding.get(
                "severity",
                ""
            )
        ).upper() == "WARNING"
    ]

    if not findings:

        paragraph = document.add_paragraph()

        paragraph.add_run(
            "No WARNING findings were identified "
            "during the assessment."
        )

        return

    table = document.add_table(
        rows=1,
        cols=4
    )

    table.style = "Table Grid"

    headers = [
        "Server",
        "Domain",
        "Finding",
        "Recommendation",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "FFF2CC"
        )

    for finding in findings:

        row = table.add_row().cells

        set_cell_text(
            row[0],
            finding.get(
                "_hostname",
                "UNKNOWN"
            )
        )

        set_cell_text(
            row[1],
            CATEGORY_NAMES.get(
                finding.get(
                    "category",
                    "Unknown"
                ),
                finding.get(
                    "category",
                    "Unknown"
                )
            )
        )

        set_cell_text(
            row[2],
            finding.get(
                "message",
                "N/A"
            )
        )

        set_cell_text(
            row[3],
            finding.get(
                "recommendation",
                "Review the finding and investigate."
            )
        )

    add_table_borders(table)


# ============================================================
# COMMON FINDINGS
# ============================================================

def add_common_findings(document, servers):

    add_section_heading(
        document,
        "7. Common Findings",
        level=1
    )

    findings = get_all_findings(servers)

    non_ok = [
        finding
        for finding in findings
        if str(
            finding.get(
                "severity",
                ""
            )
        ).upper() in [
            "CRITICAL",
            "WARNING",
            "INFO",
        ]
    ]

    finding_counter = Counter()

    for finding in non_ok:

        finding_id = finding.get(
            "id",
            finding.get(
                "message",
                "Unknown finding"
            )
        )

        finding_counter[finding_id] += 1

    if not finding_counter:

        document.add_paragraph(
            "No non-OK findings were identified."
        )

        return

    table = document.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    headers = [
        "Finding",
        "Occurrences",
        "Severity",
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True
        )

        shade_cell(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    severity_map = {}

    for finding in non_ok:

        finding_id = finding.get(
            "id",
            finding.get(
                "message",
                "Unknown finding"
            )
        )

        severity = str(
            finding.get(
                "severity",
                "UNKNOWN"
            )
        ).upper()

        severity_map[finding_id] = severity

    for finding_id, count in finding_counter.most_common():

        row = table.add_row().cells

        set_cell_text(
            row[0],
            finding_id
        )

        set_cell_text(
            row[1],
            count
        )

        set_cell_text(
            row[2],
            severity_map.get(
                finding_id,
                "UNKNOWN"
            )
        )

    add_table_borders(table)


# ============================================================
# CONSOLIDATED RECOMMENDATIONS
# ============================================================

def add_recommendations(document, servers):

    add_section_heading(
        document,
        "8. Consolidated Recommendations",
        level=1
    )

    findings = get_all_findings(servers)

    recommendations = []

    for finding in findings:

        severity = str(
            finding.get(
                "severity",
                ""
            )
        ).upper()

        if severity not in [
            "CRITICAL",
            "WARNING",
        ]:
            continue

        recommendation = finding.get(
            "recommendation"
        )

        if recommendation:

            recommendations.append(
                (
                    finding.get(
                        "_hostname",
                        "UNKNOWN"
                    ),
                    finding.get(
                        "category",
                        "Unknown"
                    ),
                    severity,
                    recommendation,
                )
            )

    if not recommendations:

        document.add_paragraph(
            "No corrective recommendations are required "
            "based on the current assessment results."
        )

        return

    seen = set()

    for hostname, category, severity, recommendation in recommendations:

        key = (
            category,
            severity,
            recommendation
        )

        if key in seen:
            continue

        seen.add(key)

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.add_run(
            f"[{severity}] "
        ).bold = True

        paragraph.add_run(
            f"{CATEGORY_NAMES.get(category, category)}: "
        ).bold = True

        paragraph.add_run(
            recommendation
        )


# ============================================================
# CONCLUSION
# ============================================================

def add_conclusion(document, servers):

    add_section_heading(
        document,
        "9. Conclusion",
        level=1
    )

    status_counts = get_status_counts(servers)

    critical = status_counts.get(
        "CRITICAL",
        0
    )

    warning = status_counts.get(
        "WARNING",
        0
    )

    healthy = status_counts.get(
        "OK",
        0
    )

    total = len(servers)

    if critical > 0:

        conclusion = (
            f"The quarterly assessment identified "
            f"{critical} server(s) with critical health findings. "
            f"These systems should be prioritized for immediate "
            f"investigation and corrective action."
        )

    elif warning > 0:

        conclusion = (
            f"The quarterly assessment identified "
            f"{warning} server(s) with warning-level findings. "
            f"No critical server-level health status was identified. "
            f"The warning findings should be reviewed and addressed "
            f"through the normal operational and change-management process."
        )

    else:

        conclusion = (
            f"The quarterly assessment covered {total} server(s), "
            f"with no server-level warning or critical status identified. "
            f"The assessed environment is operating within the configured "
            f"health-check thresholds."
        )

    document.add_paragraph(conclusion)

    document.add_paragraph()

    document.add_paragraph(
        f"Healthy servers: {healthy}/{total}"
    )

    document.add_paragraph(
        f"Warning servers: {warning}/{total}"
    )

    document.add_paragraph(
        f"Critical servers: {critical}/{total}"
    )


# ============================================================
# MAIN
# ============================================================

def generate_master_report():

    print("=" * 70)
    print("STC RHEL OS HEALTH CHECK")
    print("Master DOCX Report Generator")
    print("=" * 70)

    print()
    print(f"Reporting period : {REPORT_PERIOD}")
    print(f"Data directory   : {DATA_DIR}")
    print(f"Output file      : {OUTPUT_FILE}")
    print()

    servers = load_json_files()

    print(
        f"JSON reports loaded: {len(servers)}"
    )

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document()

    # Page setup
    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    # Default font
    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    add_title(document)

    document.add_page_break()

    add_executive_summary(
        document,
        servers
    )

    add_server_summary(
        document,
        servers
    )

    add_finding_summary(
        document,
        servers
    )

    add_domain_summary(
        document,
        servers
    )

    add_critical_findings(
        document,
        servers
    )

    add_warning_findings(
        document,
        servers
    )

    add_common_findings(
        document,
        servers
    )

    add_recommendations(
        document,
        servers
    )

    add_conclusion(
        document,
        servers
    )

    document.save(
        OUTPUT_FILE
    )

    print()
    print("=" * 70)
    print("MASTER REPORT GENERATED SUCCESSFULLY")
    print("=" * 70)

    print()
    print(f"Servers included : {len(servers)}")
    print(f"Output file      : {OUTPUT_FILE}")
    print()


def main():

    try:
        generate_master_report()

    except KeyboardInterrupt:

        print("\nOperation cancelled.")
        sys.exit(1)

    except Exception as exc:

        print()
        print("ERROR: Master report generation failed.")
        print(str(exc))

        sys.exit(1)


if __name__ == "__main__":
    main()
