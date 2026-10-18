#!/usr/bin/env python3

import json
import sys
from pathlib import Path
from collections import OrderedDict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt


# ---------------------------------------------------------
# Configuration
# ---------------------------------------------------------

CATEGORY_ORDER = [
    "cpu",
    "memory",
    "filesystem",
    "services",
    "kernel",
    "security",
    "logs",
    "Network",
    "NTP",
    "OS Version",
    "Uptime",
    "SAR",
]

CATEGORY_NAMES = {
    "cpu": "CPU Health",
    "memory": "Memory Health",
    "filesystem": "Filesystem Health",
    "services": "Services",
    "kernel": "Kernel",
    "security": "Security",
    "logs": "Logs",
    "Network": "Network",
    "NTP": "NTP",
    "OS Version": "OS Version",
    "Uptime": "Uptime",
    "SAR": "SAR",
}


# ---------------------------------------------------------
# JSON
# ---------------------------------------------------------

def load_json(json_file):
    with open(json_file, "r", encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------
# Document helpers
# ---------------------------------------------------------

def add_title(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("STC RHEL OS HEALTH CHECK REPORT")
    run.bold = True
    run.font.size = Pt(20)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("Quarterly RHEL Operating System Health Assessment")
    run.italic = True
    run.font.size = Pt(11)


def add_heading(document, text, level=1):
    document.add_heading(text, level=level)


def add_label_value(document, label, value):
    paragraph = document.add_paragraph()

    run = paragraph.add_run(f"{label}: ")
    run.bold = True

    paragraph.add_run(str(value))


# ---------------------------------------------------------
# Server Information
# ---------------------------------------------------------

def add_server_information(document, server):

    add_heading(document, "1. Server Information", 1)

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    table.rows[0].cells[0].text = "Property"
    table.rows[0].cells[1].text = "Value"

    fields = [
        ("Hostname", server.get("hostname", "N/A")),
        ("Inventory Hostname", server.get("inventory_hostname", "N/A")),
        ("FQDN", server.get("fqdn", "N/A")),
        ("IP Address", server.get("ip_address", "N/A")),
        ("System UUID (SMBIOS/DMI)", server.get("system_uuid", "N/A")),
        ("Latest Package Transaction", server.get("latest_package_transaction_date", "N/A")),
        ("Transaction Command", server.get("latest_package_transaction_command", "N/A")),
        ("Operating System", server.get("os", "N/A")),
        ("OS Version", server.get("os_version", "N/A")),
        ("Kernel", server.get("kernel", "N/A")),
        ("Architecture", server.get("architecture", "N/A")),
        ("Uptime (Days)", server.get("uptime_days", "N/A")),
    ]

    for label, value in fields:

        cells = table.add_row().cells

        cells[0].text = label
        cells[1].text = str(value)


# ---------------------------------------------------------
# Overall Health
# ---------------------------------------------------------

def add_overall_health(document, health):

    add_heading(document, "2. Overall Health Summary", 1)

    status = health.get("overall_status", "UNKNOWN")

    add_label_value(document, "Overall Health", status)

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    table.rows[0].cells[0].text = "Severity"
    table.rows[0].cells[1].text = "Count"

    for severity, key in [
        ("CRITICAL", "critical"),
        ("WARNING", "warning"),
        ("INFO", "info"),
        ("OK", "ok"),
        ("TOTAL", "total"),
    ]:

        cells = table.add_row().cells

        cells[0].text = severity
        cells[1].text = str(health.get(key, 0))


# ---------------------------------------------------------
# Group findings by category
# ---------------------------------------------------------

def group_findings(findings):

    grouped = OrderedDict()

    for finding in findings:

        category = finding.get("category", "General")

        grouped.setdefault(category, [])

        grouped[category].append(finding)

    return grouped


# ---------------------------------------------------------
# Determine domain status
# ---------------------------------------------------------

def domain_status(findings):

    severities = [
        finding.get("severity", "INFO")
        for finding in findings
    ]

    if "CRITICAL" in severities:
        return "CRITICAL"

    if "WARNING" in severities:
        return "WARNING"

    if "INFO" in severities:
        return "INFO"

    return "OK"


# ---------------------------------------------------------
# Domain Summary
# ---------------------------------------------------------

def add_domain_summary(document, grouped):

    add_heading(document, "3. Health Domain Summary", 1)

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    table.rows[0].cells[0].text = "Health Domain"
    table.rows[0].cells[1].text = "Status"

    processed = set()

    # Follow preferred domain order
    for category in CATEGORY_ORDER:

        if category not in grouped:
            continue

        status = domain_status(grouped[category])

        name = CATEGORY_NAMES.get(
            category,
            category
        )

        cells = table.add_row().cells

        cells[0].text = name
        cells[1].text = status

        processed.add(category)

    # Add any categories not in our predefined order
    for category, findings in grouped.items():

        if category in processed:
            continue

        cells = table.add_row().cells

        cells[0].text = CATEGORY_NAMES.get(
            category,
            category
        )

        cells[1].text = domain_status(findings)


# ---------------------------------------------------------
# Extract useful metric information
# ---------------------------------------------------------

def add_domain_metrics(document, findings):

    metrics = []

    for finding in findings:

        message = finding.get("message", "")

        category = finding.get("category", "").lower()

        # CPU
        if category == "cpu":

            if "utilization" in message.lower():
                metrics.append(("CPU Utilization", message))

            elif "normalized" in message.lower():
                metrics.append(("CPU Load", message))

            elif "steal" in message.lower():
                metrics.append(("CPU Steal", message))

            elif "i/o wait" in message.lower():
                metrics.append(("I/O Wait", message))

        # Memory
        elif category == "memory":

            if "memory utilization" in message.lower():
                metrics.append(("Memory Utilization", message))

            elif "swap utilization" in message.lower():
                metrics.append(("Swap Utilization", message))

            elif "swap activity" in message.lower():
                metrics.append(("Swap Activity", message))

        # Filesystem
        elif category == "filesystem":

            mountpoint = finding.get("mountpoint", "N/A")
            usage = finding.get("usage_percent", "N/A")

            metrics.append(
                (str(mountpoint), f"{usage}%")
            )

        # Services
        elif category == "services":

            if "failed" in message.lower():

                metrics.append(
                    ("Service Status", message)
                )

    if not metrics:
        return

    paragraph = document.add_paragraph()

    run = paragraph.add_run("Key Metrics")
    run.bold = True

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"

    for metric, value in metrics:

        cells = table.add_row().cells

        cells[0].text = metric
        cells[1].text = str(value)


# ---------------------------------------------------------
# Detailed Finding
# ---------------------------------------------------------

def add_detailed_finding(document, finding):

    severity = finding.get("severity", "UNKNOWN")

    message = finding.get(
        "message",
        "No message provided."
    )

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        f"{severity}: {message}"
    )

    run.bold = True

    # Evidence
    evidence = finding.get("evidence")

    if evidence:

        paragraph = document.add_paragraph()

        run = paragraph.add_run("Evidence")
        run.bold = True

        document.add_paragraph(
            str(evidence)
        )

    # Impact
    impact = finding.get("impact")

    if impact:

        paragraph = document.add_paragraph()

        run = paragraph.add_run("Impact")
        run.bold = True

        document.add_paragraph(
            str(impact)
        )

    # Recommendation
    recommendation = finding.get("recommendation")

    if recommendation:

        paragraph = document.add_paragraph()

        run = paragraph.add_run("Recommendation")
        run.bold = True

        document.add_paragraph(
            str(recommendation)
        )

    # Diagnostics
    diagnostics = finding.get(
        "diagnostic_commands",
        []
    )

    if diagnostics:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "Diagnostic Commands"
        )
        run.bold = True

        for command in diagnostics:

            document.add_paragraph(
                str(command),
                style="List Bullet"
            )

    # Root cause
    root_cause = finding.get(
        "root_cause_checks",
        []
    )

    if root_cause:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "Root Cause Checks"
        )
        run.bold = True

        for check in root_cause:

            document.add_paragraph(
                str(check),
                style="List Bullet"
            )

    # Remediation
    remediation = finding.get(
        "remediation",
        []
    )

    if remediation:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "Remediation"
        )
        run.bold = True

        for action in remediation:

            document.add_paragraph(
                str(action),
                style="List Bullet"
            )


# ---------------------------------------------------------
# Domain Section
# ---------------------------------------------------------

def add_domain_section(document, category, findings):

    name = CATEGORY_NAMES.get(
        category,
        category
    )

    status = domain_status(findings)

    add_heading(
        document,
        name,
        1
    )

    add_label_value(
        document,
        "Status",
        status
    )

    # Metrics for operational domains
    if category.lower() in [
        "cpu",
        "memory",
        "filesystem",
        "services",
    ]:

        add_domain_metrics(
            document,
            findings
        )

    # Separate findings
    warnings = []
    criticals = []
    infos = []
    oks = []

    for finding in findings:

        severity = finding.get(
            "severity",
            "INFO"
        )

        if severity == "CRITICAL":
            criticals.append(finding)

        elif severity == "WARNING":
            warnings.append(finding)

        elif severity == "INFO":
            infos.append(finding)

        else:
            oks.append(finding)

    # Important findings first
    important = (
        criticals +
        warnings +
        infos
    )

    if important:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "Findings Requiring Review"
        )

        run.bold = True

        for finding in important:

            add_detailed_finding(
                document,
                finding
            )

    # Healthy findings summarized
    if oks:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            f"Healthy Checks: {len(oks)}"
        )

        run.bold = True

        # Don't dump all OK messages.
        document.add_paragraph(
            "The checks in this domain are within "
            "the configured health thresholds."
        )


# ---------------------------------------------------------
# Final Recommendations
# ---------------------------------------------------------

def add_final_recommendations(document, findings):

    add_heading(
        document,
        "16. Final Recommendations",
        1
    )

    important = [
        finding
        for finding in findings
        if finding.get("severity") in [
            "CRITICAL",
            "WARNING"
        ]
    ]

    if not important:

        document.add_paragraph(
            "No critical or warning findings were identified. "
            "Continue routine operational monitoring."
        )

        return

    number = 1

    for finding in important:

        recommendation = finding.get(
            "recommendation"
        )

        if not recommendation:
            continue

        category = finding.get(
            "category",
            "General"
        )

        message = finding.get(
            "message",
            "Finding"
        )

        paragraph = document.add_paragraph(
            style="List Number"
        )

        paragraph.add_run(
            f"[{category}] {message} — "
        ).bold = True

        paragraph.add_run(
            str(recommendation)
        )

        number += 1


# ---------------------------------------------------------
# Conclusion
# ---------------------------------------------------------

def add_conclusion(document, health):

    add_heading(
        document,
        "17. Conclusion",
        1
    )

    status = health.get(
        "overall_status",
        "UNKNOWN"
    )

    critical = health.get(
        "critical",
        0
    )

    warning = health.get(
        "warning",
        0
    )

    if status == "CRITICAL":

        message = (
            f"The server requires immediate attention. "
            f"{critical} critical finding(s) were identified."
        )

    elif status == "WARNING":

        message = (
            f"The server health assessment resulted in "
            f"WARNING status. {warning} warning finding(s) "
            f"require operational review."
        )

    elif status == "OK":

        message = (
            "The server passed the configured health checks. "
            "No critical or warning findings were identified."
        )

    else:

        message = (
            "The server health status could not be determined."
        )

    document.add_paragraph(message)


# ---------------------------------------------------------
# Main report generation
# ---------------------------------------------------------

def generate_report(json_file, output_file):

    data = load_json(json_file)

    document = Document()

    # Page margins
    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Default font
    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    add_title(document)

    document.add_paragraph()

    server = data.get(
        "server",
        {}
    )

    health = data.get(
        "health",
        {}
    )

    findings = data.get(
        "findings",
        []
    )

    grouped = group_findings(
        findings
    )

    # Sections
    add_server_information(
        document,
        server
    )

    document.add_paragraph()

    add_overall_health(
        document,
        health
    )

    document.add_paragraph()

    add_domain_summary(
        document,
        grouped
    )

    document.add_page_break()

    # Domain sections
    section_number = 4

    for category in CATEGORY_ORDER:

        if category not in grouped:
            continue

        add_domain_section(
            document,
            category,
            grouped[category]
        )

        document.add_paragraph()

        section_number += 1

    # Any unknown categories
    for category, category_findings in grouped.items():

        if category in CATEGORY_ORDER:
            continue

        add_domain_section(
            document,
            category,
            category_findings
        )

        document.add_paragraph()

    add_final_recommendations(
        document,
        findings
    )

    add_conclusion(
        document,
        health
    )

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document.save(
        output_file
    )


# ---------------------------------------------------------
# CLI
# ---------------------------------------------------------

def main():

    if len(sys.argv) != 3:

        print(
            "Usage: python3 generate_server_report.py "
            "<input.json> <output.docx>"
        )

        sys.exit(1)

    json_file = Path(
        sys.argv[1]
    )

    output_file = Path(
        sys.argv[2]
    )

    if not json_file.exists():

        print(
            f"ERROR: JSON file not found: {json_file}"
        )

        sys.exit(1)

    generate_report(
        json_file,
        output_file
    )

    print(
        "Report generated successfully:"
    )

    print(
        output_file
    )


if __name__ == "__main__":
    main()
