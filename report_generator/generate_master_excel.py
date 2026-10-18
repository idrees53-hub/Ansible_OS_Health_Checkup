#!/usr/bin/env python3

from pathlib import Path
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import json
import sys


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

REPORT_PERIOD = "2026-Q4"

DATA_DIR = (
    PROJECT_ROOT
    / "reports"
    / REPORT_PERIOD
    / "data"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "reports"
    / REPORT_PERIOD
    / "STC_OS_Health_Summary.xlsx"
)


CATEGORY_NAMES = {
    "cpu": "CPU",
    "memory": "Memory",
    "filesystem": "Filesystem",
    "services": "Services",
    "kernel": "Kernel",
    "security": "Security",
    "logs": "Logs",
    "Network": "Network",
    "NTP": "NTP / Time Synchronization",
    "OS Version": "OS Version",
    "Uptime": "Uptime",
    "SAR": "SAR / Historical Performance",
}


# ============================================================
# STYLES
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF"
)

TITLE_FONT = Font(
    bold=True,
    size=16
)

SUBTITLE_FONT = Font(
    bold=True,
    size=11
)

THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)


# ============================================================
# DATA LOADING
# ============================================================

def load_json_files():

    if not DATA_DIR.exists():

        print("ERROR: Data directory does not exist:")
        print(DATA_DIR)

        sys.exit(1)

    json_files = sorted(
        DATA_DIR.glob("*.json")
    )

    if not json_files:

        print("ERROR: No JSON report files found.")
        print(DATA_DIR)

        sys.exit(1)

    servers = []

    for json_file in json_files:

        try:

            with open(
                json_file,
                "r",
                encoding="utf-8"
            ) as file:

                data = json.load(file)

            servers.append(data)

        except Exception as exc:

            print(
                f"WARNING: Could not read "
                f"{json_file.name}: {exc}"
            )

    if not servers:

        print(
            "ERROR: No valid JSON reports could be loaded."
        )

        sys.exit(1)

    return servers


# ============================================================
# WORKSHEET HELPERS
# ============================================================

def style_header(row):

    for cell in row:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = THIN_BORDER


def style_all_cells(ws):

    for row in ws.iter_rows():

        for cell in row:

            cell.border = THIN_BORDER

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


def auto_width(ws, maximum=45):

    for column_cells in ws.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:

            if cell.value is not None:

                length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    length
                )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            maximum
        )


def freeze_header(ws):

    ws.freeze_panes = "A2"


def add_filters(ws):

    if ws.max_row >= 2:

        ws.auto_filter.ref = ws.dimensions


def set_status_format(cell):

    status = str(
        cell.value or ""
    ).upper()

    if status == "CRITICAL":

        cell.fill = PatternFill(
            "solid",
            fgColor="F4CCCC"
        )

    elif status == "WARNING":

        cell.fill = PatternFill(
            "solid",
            fgColor="FFF2CC"
        )

    elif status == "OK":

        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAD3"
        )

    elif status == "INFO":

        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )


# ============================================================
# SHEET 1 - DASHBOARD
# ============================================================

def create_dashboard(wb, servers):

    ws = wb.create_sheet(
        "Dashboard",
        0
    )

    ws["A1"] = (
        "STC RHEL OS HEALTH CHECK"
    )

    ws["A1"].font = TITLE_FONT

    ws["A2"] = (
        f"Quarterly Summary - {REPORT_PERIOD}"
    )

    ws["A2"].font = SUBTITLE_FONT

    status_counts = Counter()

    finding_counts = Counter()

    for server in servers:

        health = server.get(
            "health",
            {}
        )

        status = str(
            health.get(
                "overall_status",
                "UNKNOWN"
            )
        ).upper()

        status_counts[status] += 1

        finding_counts[
            "CRITICAL"
        ] += int(
            health.get(
                "critical",
                0
            )
        )

        finding_counts[
            "WARNING"
        ] += int(
            health.get(
                "warning",
                0
            )
        )

        finding_counts[
            "INFO"
        ] += int(
            health.get(
                "info",
                0
            )
        )

        finding_counts[
            "OK"
        ] += int(
            health.get(
                "ok",
                0
            )
        )

    total_servers = len(servers)

    # --------------------------------------------------------
    # Server Status
    # --------------------------------------------------------

    ws["A4"] = "SERVER STATUS"
    ws["A4"].font = SUBTITLE_FONT

    dashboard_rows = [
        [
            "Total Servers",
            total_servers
        ],
        [
            "Healthy Servers",
            status_counts.get(
                "OK",
                0
            )
        ],
        [
            "Warning Servers",
            status_counts.get(
                "WARNING",
                0
            )
        ],
        [
            "Critical Servers",
            status_counts.get(
                "CRITICAL",
                0
            )
        ],
    ]

    row_number = 5

    for label, value in dashboard_rows:

        ws.cell(
            row=row_number,
            column=1,
            value=label
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value
        )

        row_number += 1

    # --------------------------------------------------------
    # Finding Counts
    # --------------------------------------------------------

    ws["D4"] = "FINDING COUNTS"
    ws["D4"].font = SUBTITLE_FONT

    finding_rows = [
        [
            "CRITICAL",
            finding_counts["CRITICAL"]
        ],
        [
            "WARNING",
            finding_counts["WARNING"]
        ],
        [
            "INFO",
            finding_counts["INFO"]
        ],
        [
            "OK",
            finding_counts["OK"]
        ],
    ]

    row_number = 5

    for severity, value in finding_rows:

        ws.cell(
            row=row_number,
            column=4,
            value=severity
        )

        ws.cell(
            row=row_number,
            column=5,
            value=value
        )

        set_status_format(
            ws.cell(
                row=row_number,
                column=4
            )
        )

        row_number += 1

    # --------------------------------------------------------
    # Health Percentage
    # --------------------------------------------------------

    ws["A11"] = "HEALTH PERCENTAGE"
    ws["A11"].font = SUBTITLE_FONT

    healthy = status_counts.get(
        "OK",
        0
    )

    health_percentage = (
        (healthy / total_servers) * 100
        if total_servers
        else 0
    )

    ws["A12"] = (
        "Servers with overall OK status"
    )

    ws["B12"] = (
        health_percentage / 100
    )

    ws["B12"].number_format = "0.0%"

    # --------------------------------------------------------
    # Note
    # --------------------------------------------------------

    ws["A14"] = (
        "Note: Detailed findings are available "
        "in the Findings sheet."
    )

    ws.merge_cells(
        "A14:E14"
    )

    ws["A14"].alignment = Alignment(
        wrap_text=True
    )

    # STC dashboard charts: server status and finding severity.

    if total_servers:

        status_chart = BarChart()

        status_chart.type = "bar"

        status_chart.style = 10

        status_chart.title = "Server Health Status"

        status_chart.y_axis.title = "Servers"

        status_chart.x_axis.title = "Status"

        status_data = Reference(ws, min_col=2, min_row=6, max_row=8)

        status_categories = Reference(ws, min_col=1, min_row=6, max_row=8)

        status_chart.add_data(status_data, titles_from_data=False)

        status_chart.set_categories(status_categories)

        status_chart.height = 7

        status_chart.width = 12

        ws.add_chart(status_chart, "G4")


    if sum(finding_counts.values()):

        severity_chart = PieChart()

        severity_chart.title = "Finding Severity Distribution"

        severity_data = Reference(ws, min_col=5, min_row=4, max_row=8)

        severity_categories = Reference(ws, min_col=4, min_row=5, max_row=8)

        severity_chart.add_data(severity_data, titles_from_data=True)

        severity_chart.set_categories(severity_categories)

        severity_chart.height = 8

        severity_chart.width = 12

        ws.add_chart(severity_chart, "G20")


    style_all_cells(ws)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 5
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 18


# ============================================================
# SHEET 2 - SERVER SUMMARY
# ============================================================

def create_server_summary(wb, servers):

    # IMPORTANT:
    # Create a NEW worksheet.
    # Do not use wb.active because Dashboard
    # is already the active worksheet.

    ws = wb.create_sheet(
        "Server Summary"
    )

    headers = [
        "Server",
        "Inventory Host",
        "FQDN",
        "IP Address",
        "System UUID",
        "OS",
        "OS Version",
        "Kernel",
        "Architecture",
        "Uptime Days",
        "Overall Status",
        "Critical",
        "Warning",
        "Info",
        "OK",
        "Total Findings",
    ]

    ws.append(headers)

    style_header(
        ws[1]
    )

    for server in servers:

        info = server.get(
            "server",
            {}
        )

        health = server.get(
            "health",
            {}
        )

        row = [

            info.get(
                "hostname",
                "N/A"
            ),

            info.get(
                "inventory_hostname",
                "N/A"
            ),

            info.get(
                "fqdn",
                "N/A"
            ),

            info.get(
                "ip_address",
                "N/A"
            ),


            info.get(

                "system_uuid",

                "N/A"

            ),

            info.get(
                "os",
                "N/A"
            ),

            info.get(
                "os_version",
                "N/A"
            ),

            info.get(
                "kernel",
                "N/A"
            ),

            info.get(
                "architecture",
                "N/A"
            ),

            info.get(
                "uptime_days",
                0
            ),

            health.get(
                "overall_status",
                "UNKNOWN"
            ),

            health.get(
                "critical",
                0
            ),

            health.get(
                "warning",
                0
            ),

            health.get(
                "info",
                0
            ),

            health.get(
                "ok",
                0
            ),

            health.get(
                "total",
                0
            ),
        ]

        ws.append(row)

    for row in ws.iter_rows(
        min_row=2,
        min_col=10,
        max_col=10
    ):

        set_status_format(
            row[0]
        )

    style_all_cells(ws)
    freeze_header(ws)
    add_filters(ws)
    auto_width(ws)


# ============================================================
# SHEET 3 - FINDINGS
# ============================================================

def create_findings_sheet(wb, servers):

    ws = wb.create_sheet(
        "Findings"
    )

    headers = [
        "Server",
        "Domain",
        "Finding ID",
        "Severity",
        "Message",
        "Evidence",
        "Impact",
        "Recommendation",
        "Diagnostic Commands",
        "Root Cause Checks",
        "Remediation",
    ]

    ws.append(headers)

    style_header(
        ws[1]
    )

    for server in servers:

        hostname = server.get(
            "server",
            {}
        ).get(
            "hostname",
            "UNKNOWN"
        )

        findings = server.get(
            "findings",
            []
        )

        for finding in findings:

            diagnostics = finding.get(
                "diagnostic_commands",
                []
            )

            root_cause = finding.get(
                "root_cause_checks",
                []
            )

            remediation = finding.get(
                "remediation",
                []
            )

            if isinstance(
                diagnostics,
                list
            ):

                diagnostics = "\n".join(
                    str(x)
                    for x in diagnostics
                )

            if isinstance(
                root_cause,
                list
            ):

                root_cause = "\n".join(
                    str(x)
                    for x in root_cause
                )

            if isinstance(
                remediation,
                list
            ):

                remediation = "\n".join(
                    str(x)
                    for x in remediation
                )

            row = [

                hostname,

                CATEGORY_NAMES.get(
                    finding.get(
                        "category",
                        "Unknown"
                    ),
                    finding.get(
                        "category",
                        "Unknown"
                    )
                ),

                finding.get(
                    "id",
                    "N/A"
                ),

                finding.get(
                    "severity",
                    "UNKNOWN"
                ),

                finding.get(
                    "message",
                    "N/A"
                ),

                finding.get(
                    "evidence",
                    "N/A"
                ),

                finding.get(
                    "impact",
                    "N/A"
                ),

                finding.get(
                    "recommendation",
                    "N/A"
                ),

                diagnostics,

                root_cause,

                remediation,
            ]

            ws.append(row)

    for row in ws.iter_rows(
        min_row=2,
        min_col=4,
        max_col=4
    ):

        set_status_format(
            row[0]
        )

    style_all_cells(ws)
    freeze_header(ws)
    add_filters(ws)
    auto_width(
        ws,
        maximum=55
    )


# ============================================================
# SHEET 4 - DOMAIN SUMMARY
# ============================================================

def create_domain_summary(wb, servers):

    ws = wb.create_sheet(
        "Domain Summary"
    )

    headers = [
        "Health Domain",
        "Critical",
        "Warning",
        "Info",
        "OK",
        "Total",
    ]

    ws.append(headers)

    style_header(
        ws[1]
    )

    domain_counts = defaultdict(
        Counter
    )

    for server in servers:

        for finding in server.get(
            "findings",
            []
        ):

            category = finding.get(
                "category",
                "Unknown"
            )

            severity = str(
                finding.get(
                    "severity",
                    "UNKNOWN"
                )
            ).upper()

            domain_counts[
                category
            ][severity] += 1

    for category in domain_counts:

        counts = domain_counts[
            category
        ]

        critical = counts.get(
            "CRITICAL",
            0
        )

        warning = counts.get(
            "WARNING",
            0
        )

        info = counts.get(
            "INFO",
            0
        )

        ok = counts.get(
            "OK",
            0
        )

        total = (
            critical
            + warning
            + info
            + ok
        )

        ws.append(
            [
                CATEGORY_NAMES.get(
                    category,
                    category
                ),
                critical,
                warning,
                info,
                ok,
                total,
            ]
        )

    style_all_cells(ws)
    freeze_header(ws)
    add_filters(ws)
    auto_width(ws)


# ============================================================
# SHEET 5 - CRITICAL SERVERS
# ============================================================

def create_critical_servers(wb, servers):

    ws = wb.create_sheet(
        "Critical Servers"
    )

    headers = [
        "Server",
        "OS",
        "Critical Findings",
        "Warning Findings",
        "Overall Status",
    ]

    ws.append(headers)

    style_header(
        ws[1]
    )

    for server in servers:

        health = server.get(
            "health",
            {}
        )

        critical = int(
            health.get(
                "critical",
                0
            )
        )

        if critical <= 0:
            continue

        info = server.get(
            "server",
            {}
        )

        row = [

            info.get(
                "hostname",
                "UNKNOWN"
            ),

            f"{info.get('os', 'N/A')} "
            f"{info.get('os_version', '')}",

            critical,

            health.get(
                "warning",
                0
            ),

            health.get(
                "overall_status",
                "UNKNOWN"
            ),
        ]

        ws.append(row)

    if ws.max_row == 1:

        ws.append(
            [
                "No critical servers identified",
                "",
                0,
                0,
                "OK",
            ]
        )

    style_all_cells(ws)
    freeze_header(ws)
    add_filters(ws)
    auto_width(ws)


# ============================================================
# SHEET 6 - WARNING SERVERS
# ============================================================

def create_warning_servers(wb, servers):

    ws = wb.create_sheet(
        "Warning Servers"
    )

    headers = [
        "Server",
        "OS",
        "Critical Findings",
        "Warning Findings",
        "Overall Status",
    ]

    ws.append(headers)

    style_header(
        ws[1]
    )

    for server in servers:

        health = server.get(
            "health",
            {}
        )

        critical = int(
            health.get(
                "critical",
                0
            )
        )

        warning = int(
            health.get(
                "warning",
                0
            )
        )

        if warning <= 0:
            continue

        info = server.get(
            "server",
            {}
        )

        row = [

            info.get(
                "hostname",
                "UNKNOWN"
            ),

            f"{info.get('os', 'N/A')} "
            f"{info.get('os_version', '')}",

            critical,

            warning,

            health.get(
                "overall_status",
                "UNKNOWN"
            ),
        ]

        ws.append(row)

    if ws.max_row == 1:

        ws.append(
            [
                "No warning servers identified",
                "",
                0,
                0,
                "OK",
            ]
        )

    style_all_cells(ws)
    freeze_header(ws)
    add_filters(ws)
    auto_width(ws)


# ============================================================
# MAIN
# ============================================================

def generate_excel_report():

    print("=" * 70)
    print("STC RHEL OS HEALTH CHECK")
    print("Master Excel Report Generator")
    print("=" * 70)

    print()

    print(
        f"Reporting period : {REPORT_PERIOD}"
    )

    print(
        f"Data directory   : {DATA_DIR}"
    )

    print(
        f"Output file      : {OUTPUT_FILE}"
    )

    print()

    servers = load_json_files()

    print(
        f"JSON reports loaded: {len(servers)}"
    )

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook = Workbook()

    # Remove the default empty worksheet.
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # Create sheets in the required order.

    create_dashboard(
        workbook,
        servers
    )

    create_server_summary(
        workbook,
        servers
    )

    create_findings_sheet(
        workbook,
        servers
    )

    create_domain_summary(
        workbook,
        servers
    )

    create_critical_servers(
        workbook,
        servers
    )

    create_warning_servers(
        workbook,
        servers
    )

    workbook.save(
        OUTPUT_FILE
    )

    print()

    print("=" * 70)
    print("MASTER EXCEL REPORT GENERATED SUCCESSFULLY")
    print("=" * 70)

    print()

    print(
        f"Servers included : {len(servers)}"
    )

    print(
        f"Output file      : {OUTPUT_FILE}"
    )

    print()


# ============================================================
# PROGRAM ENTRY POINT
# ============================================================

def main():

    try:

        generate_excel_report()

    except KeyboardInterrupt:

        print(
            "\nOperation cancelled."
        )

        sys.exit(1)

    except Exception as exc:

        print()

        print(
            "ERROR: Excel report generation failed."
        )

        print(
            str(exc)
        )

        sys.exit(1)


if __name__ == "__main__":

    main()
